import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def create_client_proposal():
    doc = Document()

    # Set Margins to 0.75 in
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Primary colors
    NAVY = RGBColor(15, 41, 66)       # #0F2942
    STEEL = RGBColor(30, 58, 138)     # #1E3A8A
    CHARCOAL = RGBColor(30, 41, 59)   # #1E293B
    MUTED = RGBColor(100, 116, 139)   # #64748B

    # Helper styling functions
    def add_title(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = NAVY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.font.color.rgb = MUTED
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(12)
        return p

    def add_heading1(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = NAVY
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        return p

    def add_heading2(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = STEEL
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        return p

    def add_body(text, bold_prefix=None, italic=False):
        p = doc.add_paragraph()
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.font.name = 'Calibri'
            run_b.font.size = Pt(10)
            run_b.font.bold = True
            run_b.font.color.rgb = CHARCOAL
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
        run.font.italic = italic
        run.font.color.rgb = CHARCOAL
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        return p

    def add_bullet(title, text):
        p = doc.add_paragraph(style='List Bullet')
        run_t = p.add_run(title + ": ")
        run_t.font.name = 'Calibri'
        run_t.font.size = Pt(10)
        run_t.font.bold = True
        run_t.font.color.rgb = CHARCOAL

        run_d = p.add_run(text)
        run_d.font.name = 'Calibri'
        run_d.font.size = Pt(10)
        run_d.font.color.rgb = CHARCOAL
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        return p

    # -------------------------------------------------------------------------
    # HEADER & DOCUMENT INFO BLOCK
    # -------------------------------------------------------------------------
    add_title("COMMERCIAL PROPOSAL & SYSTEM IMPLEMENTATION PLAN")
    add_subtitle("Venture Visionary Visa & Business Consultancy Management Platform (v2.0)")

    # Overview Table
    info_table = doc.add_table(rows=2, cols=3)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(info_table, color="CBD5E1", sz="4")

    info_data = [
        [("Client Name:", " Venture Visionary Consultancy"), ("Prepared For:", " Executive Management, Qatar"), ("Date:", " October 2026")],
        [("System Version:", " Version 2.0 Management Platform"), ("Delivery Model:", " Incremental Phased Delivery"), ("Target Market:", " Qatar & Global Operations")]
    ]

    for r_idx, row in enumerate(info_data):
        for c_idx, (label, val) in enumerate(row):
            cell = info_table.cell(r_idx, c_idx)
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run_lbl = p.add_run(label)
            run_lbl.font.name = 'Calibri'
            run_lbl.font.size = Pt(9)
            run_lbl.font.bold = True
            run_lbl.font.color.rgb = NAVY
            run_v = p.add_run(val)
            run_v.font.name = 'Calibri'
            run_v.font.size = Pt(9)
            run_v.font.color.rgb = CHARCOAL

    # Set column widths
    for row in info_table.rows:
        row.cells[0].width = Inches(2.3)
        row.cells[1].width = Inches(2.6)
        row.cells[2].width = Inches(2.1)

    # -------------------------------------------------------------------------
    # 1. EXECUTIVE SUMMARY (PURE BUSINESS VALUE)
    # -------------------------------------------------------------------------
    add_heading1("1. EXECUTIVE SUMMARY & BUSINESS VALUE")
    add_body(
        "This commercial proposal presents the complete plan to build and deploy the centralized management platform for Venture Visionary Visa and Business Consultancy. "
        "Unlike generic software or basic contact managers, this platform is specifically tailored to run every operational step of your visa consultancy and business setup operations in Qatar."
    )
    add_body(
        "The system solves the day-to-day challenges faced by your firm: eliminating lost leads, automating document checklists, preventing staff errors, stopping constant client phone inquiries with self-service tracking, and providing management with complete visibility over revenue and profit margins."
    )
    add_body(
        "Why this proposal is cost-effective: ",
        bold_prefix="Smart Development Strategy: "
    )
    add_body(
        "By utilizing battle-tested software foundations and modern AI-augmented developer workflows, we have eliminated the need to build standard infrastructure from scratch. "
        "By adopting a steady, phased development pace, total labor costs have been cut by a full 50% without compromising any security, speed, or functional capability."
    )

    # -------------------------------------------------------------------------
    # 2. WHAT THE SYSTEM DOES (NO JARGON)
    # -------------------------------------------------------------------------
    add_heading1("2. CORE SYSTEM CAPABILITIES & BENEFITS")
    add_body("The platform is structured into straightforward operational hubs designed for your team, clients, and partners:")

    add_bullet(
        "Multi-Channel Lead Capture & Unified WhatsApp Inbox",
        "Leads from Facebook Ads, Instagram, your Website, WhatsApp, and Walk-ins flow directly into one central inbox. Automatically assigns leads to sales staff (Round Robin), flags duplicate phone/QID numbers, and logs all communication history so no customer is forgotten."
    )
    add_bullet(
        "193-Country Visa & Business Setup Master Directory",
        "A central knowledge base for embassy requirements, official visa fees, VAC locations, processing times, and Qatar resident rules. Your research team can update rules with management review and approval before publishing."
    )
    add_bullet(
        "Instant Dynamic Document Checklists & Expiry Alerts",
        "Select the destination, nationality, Qatar residency, and travel purpose, and the system instantly generates the exact document requirements for the applicant. Alerts staff automatically before passports, QIDs, or bank statements expire."
    )
    add_bullet(
        "End-to-End Case Progression & Quality Control (QC)",
        "Guides cases through each stage: Onboarding → Document Verification → Application Ready → Submitted → Biometrics → Under Processing → Decision. Staff can review and mark documents as Verified or Rejected with reasons, preventing embassy rejections."
    )
    add_bullet(
        "Dedicated Client Tracking Portal",
        "Clients log in securely with their mobile number (OTP) to view their application progress timeline, upload requested documents directly from their phone, and download approved visas and receipts. Internal staff notes remain 100% private."
    )
    add_bullet(
        "B2B Partner & Sub-Agent Portal",
        "Travel agencies and corporate partners get their own login to submit client cases, access negotiated wholesale net rates, track files, and manage their credit limit or pre-paid wallet balance."
    )
    add_bullet(
        "Financial Management, Quotations & Profit Margins",
        "Quickly generate professional itemized quotations and invoices showing Government fees, Embassy fees, and your consultancy fee. Automatically calculates real gross profit margins per case and manages refunds with manager approval."
    )
    add_bullet(
        "Executive Management Dashboard",
        "Gives leadership real-time visibility over daily sales, active cases, revenue, overdue tasks, pending documents, and staff workload with one-click Excel reports."
    )

    # -------------------------------------------------------------------------
    # 3. INVESTMENT OPTIONS (CLEAR PRICING)
    # -------------------------------------------------------------------------
    add_heading1("3. COMMERCIAL INVESTMENT & PRICING")
    add_body("To give you maximum financial flexibility, we provide two transparent pricing options:")

    add_heading2("OPTION 1: Pure Software Development (50% Labor Reduced)")
    add_body("Covers 100% of software engineering, testing, customization, and deployment. You pay for your own hosting and messaging services directly with providers.")

    opt1_table = doc.add_table(rows=4, cols=4)
    opt1_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(opt1_table, color="CBD5E1", sz="4")

    headers1 = ["Project Milestone", "Scope Included", "Timeline", "Investment"]
    for idx, text in enumerate(headers1):
        cell = opt1_table.cell(0, idx)
        set_cell_background(cell, "0F2942")
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    opt1_rows = [
        (
            "Milestone 1:\nCore Operations MVP\n(Phase 1 + 2)",
            "• Staff CRM, Lead Engine & WhatsApp Unified Inbox\n• 193-Country Visa Directory & Governance Rules\n• Dynamic Document Checklists & Expiry Warnings\n• Document Quality Control (QC) & Case Workflows\n• Executive Dashboard & Full Audit History",
            "~2 Months\n(Steady Pace)",
            "BDT 85,000 – 1,12,500\n(~QAR 2,550 – 3,450)\n(~$700 – $950 USD)"
        ),
        (
            "Milestone 2:\nPortals & Finance\n(Phase 3 to 6)",
            "• Client Self-Service Tracking & Document Upload Portal\n• Professional Quotation, Invoicing & Profit Tracking\n• Multi-Tenant B2B Sub-Agent Portal & Wallet\n• Advanced Automations, Task Reminders & Reports",
            "~3 Months\n(Steady Pace)",
            "BDT 80,000 – 1,20,000\n(~QAR 2,400 – 3,650)\n(~$650 – $1,000 USD)"
        ),
        (
            "TOTAL FULL PLATFORM\n(Complete 6 Phases)",
            "Complete end-to-end platform matching all blueprint specifications, complete with staff training and handover.",
            "~5 Months\nTotal",
            "BDT 1,65,000 – 2,32,500\n(~QAR 5,000 – 7,100)\n(~$1,375 – $1,950 USD)"
        )
    ]

    for r_idx, row in enumerate(opt1_rows, start=1):
        for c_idx, text in enumerate(row):
            cell = opt1_table.cell(r_idx, c_idx)
            if r_idx == 3:
                set_cell_background(cell, "F1F5F9")
            else:
                set_cell_background(cell, "FFFFFF" if r_idx % 2 == 1 else "F8FAFC")
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.size = Pt(9)
            if r_idx == 3 or c_idx == 0 or c_idx == 3:
                run.font.bold = True
            run.font.color.rgb = CHARCOAL

    for row in opt1_table.rows:
        row.cells[0].width = Inches(1.8)
        row.cells[1].width = Inches(2.7)
        row.cells[2].width = Inches(1.1)
        row.cells[3].width = Inches(1.4)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_heading2("OPTION 2: All-Inclusive Turnkey Package (Development + 4 Months Hosting & Messaging)")
    add_body("Zero hassle package: We bundle custom development with 4 months of pre-funded cloud servers, automated backups, encrypted document storage, official WhatsApp Business API messaging, and SMS verification.")

    opt2_table = doc.add_table(rows=3, cols=4)
    opt2_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(opt2_table, color="CBD5E1", sz="4")

    headers2 = ["Turnkey Package Tier", "Software Scope Included", "Included 4-Month Infrastructure", "All-Inclusive Investment"]
    for idx, text in enumerate(headers2):
        cell = opt2_table.cell(0, idx)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    opt2_rows = [
        (
            "MVP Turnkey Package\n(Core MVP + 4 Mo. Infra)",
            "• Staff CRM, Lead Engine & Unified Inbox\n• 193-Country Visa Master Directory\n• Dynamic Checklists & QC Document Vault\n• SLA Alerts & Management Dashboard",
            "• 4 Months Dedicated High-Speed Cloud Server\n• 4 Months Encrypted Document Storage\n• Official WhatsApp Business API Credit\n• Client Login SMS Verification Credits",
            "BDT 1,35,000 – 2,16,500\n(~QAR 4,100 – 6,600)\n(~$1,125 – $1,800 USD)"
        ),
        (
            "Complete Enterprise Turnkey\n(Full Platform + 4 Mo. Infra)",
            "• All 3 Portals (Staff, Client, B2B Partner)\n• Quotation, Multi-Currency & Profit Ledger\n• B2B Net Rates & Partner Credit Wallets\n• Complete Workflow Automation & Reports",
            "• 4 Months Enterprise Cloud Server & Backups\n• 4 Months Encrypted Document Vault\n• Higher WhatsApp Volume Messaging Reserve\n• Transactional SMS & Email Credits",
            "BDT 2,15,000 – 3,36,500\n(~QAR 6,550 – 10,250)\n(~$1,800 – $2,800 USD)"
        )
    ]

    for r_idx, row in enumerate(opt2_rows, start=1):
        for c_idx, text in enumerate(row):
            cell = opt2_table.cell(r_idx, c_idx)
            set_cell_background(cell, "FFFFFF" if r_idx % 2 == 1 else "F8FAFC")
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.size = Pt(9)
            if c_idx == 0 or c_idx == 3:
                run.font.bold = True
            run.font.color.rgb = CHARCOAL

    for row in opt2_table.rows:
        row.cells[0].width = Inches(1.8)
        row.cells[1].width = Inches(2.2)
        row.cells[2].width = Inches(1.6)
        row.cells[3].width = Inches(1.4)

    # -------------------------------------------------------------------------
    # 4. DELIVERY SCHEDULE & IMPLEMENTATION ROADMAP
    # -------------------------------------------------------------------------
    add_heading1("4. DELIVERY SCHEDULE & ROLLOUT ROADMAP")
    add_body(
        "Development is structured into practical, low-risk business milestones so your team can begin using and benefiting from the system early:"
    )

    add_bullet(
        "Stage 1: Setup & Core Operations (Month 1 – 2)",
        "Deploy lead capture from your website and social media, set up the WhatsApp unified inbox, configure the 193-country visa master database, and launch the document QC verification engine for internal staff."
    )
    add_bullet(
        "Stage 2: Client Portal & Quotation Invoicing (Month 3 – 4)",
        "Launch client self-service tracking and secure document uploads, and activate financial quotations, invoicing, and real-time profit margin calculations."
    )
    add_bullet(
        "Stage 3: B2B Partner Portal & Final Handover (Month 4 – 5)",
        "Onboard B2B travel partners, set up net rate cards and credit wallets, conduct full team training, and perform final production handover."
    )

    # -------------------------------------------------------------------------
    # 5. PAYMENT SCHEDULE & TRANSPARENT MILESTONES
    # -------------------------------------------------------------------------
    add_heading1("5. PAYMENT MILESTONES")
    add_body("Payments are linked strictly to verified completion of milestones:")

    pay_table = doc.add_table(rows=5, cols=4)
    pay_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(pay_table, color="CBD5E1", sz="4")

    pay_headers = ["Milestone", "Deliverable Condition", "Share", "Option 1 (Dev Only)"]
    for idx, text in enumerate(pay_headers):
        cell = pay_table.cell(0, idx)
        set_cell_background(cell, "0F2942")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    pay_rows = [
        ("Milestone 1: Project Kickoff", "Signing of agreement, system setup, and approval of design blueprints.", "25%", "BDT 41,250 – 58,125"),
        ("Milestone 2: Core System Delivery", "Delivery & demonstration of Core Staff CRM, Lead Engine, Visa Directory, and Document QC.", "35%", "BDT 57,750 – 81,375"),
        ("Milestone 3: Portals & Finance", "Delivery of Client Tracking Portal, Invoicing/Quotation engine, and B2B Partner Portal.", "25%", "BDT 41,250 – 58,125"),
        ("Milestone 4: Handover & Training", "Final testing, staff training, documentation handover, and official live deployment.", "15%", "BDT 24,750 – 34,875"),
    ]

    for r_idx, row in enumerate(pay_rows, start=1):
        for c_idx, text in enumerate(row):
            cell = pay_table.cell(r_idx, c_idx)
            set_cell_background(cell, "FFFFFF" if r_idx % 2 == 1 else "F8FAFC")
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.size = Pt(9)
            if c_idx == 0 or c_idx == 2:
                run.font.bold = True
            run.font.color.rgb = CHARCOAL

    for row in pay_table.rows:
        row.cells[0].width = Inches(2.0)
        row.cells[1].width = Inches(2.9)
        row.cells[2].width = Inches(0.8)
        row.cells[3].width = Inches(1.3)

    # -------------------------------------------------------------------------
    # 6. SIGN-OFF & ACCEPTANCE
    # -------------------------------------------------------------------------
    add_heading1("6. AGREEMENT & SIGN-OFF")
    add_body("By signing below, the parties confirm their agreement to the scope, commercial terms, and milestones outlined in this proposal:")

    sign_table = doc.add_table(rows=3, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(sign_table, color="CBD5E1", sz="4")

    sign_cells = [
        [("Prepared By:", ""), ("Accepted & Approved By:", "")],
        [("Development Team Lead\nSoftware Implementation Solutions", ""), ("Executive Management / Authorized Signatory\nVenture Visionary Visa & Business Consultancy, Qatar", "")],
        [("Signature: ___________________________\n\nDate: _______________________________", ""), ("Signature: ___________________________\n\nDate: _______________________________", "")]
    ]

    for r_idx, row in enumerate(sign_cells):
        for c_idx, (text, _) in enumerate(row):
            cell = sign_table.cell(r_idx, c_idx)
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.size = Pt(9.5)
            if r_idx == 0:
                run.font.bold = True
                run.font.color.rgb = NAVY
            else:
                run.font.color.rgb = CHARCOAL

    for row in sign_table.rows:
        row.cells[0].width = Inches(3.5)
        row.cells[1].width = Inches(3.5)

    output_path = "/home/sadik/sadik/crm/Venture_Visionary_CRM_Client_Proposal.docx"
    doc.save(output_path)
    print(f"DOCX created successfully at: {output_path}")

if __name__ == '__main__':
    create_client_proposal()
